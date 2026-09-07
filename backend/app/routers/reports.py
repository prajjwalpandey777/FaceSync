import io
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.database import get_db
from app.dependencies import get_current_teacher, get_owned_class
from app.schemas import ClassReport, SessionReport, RosterEntry, OverallAttendanceRow
from app import models

router = APIRouter(prefix="/classes/{class_id}/reports", tags=["reports"])


def _build_session_report(db: Session, session: models.AttendanceSession, students) -> SessionReport:
    records = {
        r.student_id: r
        for r in db.query(models.AttendanceRecord).filter(models.AttendanceRecord.session_id == session.id)
    }
    present, absent = [], []
    for s in students:
        record = records.get(s.id)
        status_ = record.status if record else "absent"
        entry = RosterEntry(id=s.id, name=s.name, reg_no=s.reg_no, status=status_, confidence=record.confidence if record else None)
        (present if status_ == "present" else absent).append(entry)

    total = len(students)
    rate = round(len(present) / total * 100, 1) if total else 0.0
    return SessionReport(
        session_id=session.id,
        date=session.started_at,
        present=present,
        absent=absent,
        total=total,
        present_count=len(present),
        attendance_rate=rate,
    )


@router.get("", response_model=ClassReport)
def class_report(
    class_id: int,
    session_id: Optional[int] = None,
    db: Session = Depends(get_db),
    teacher: models.Teacher = Depends(get_current_teacher),
):
    school_class = get_owned_class(class_id, db, teacher)
    students = school_class.students
    sessions = (
        db.query(models.AttendanceSession)
        .filter(models.AttendanceSession.class_id == class_id, models.AttendanceSession.status == "ended")
        .order_by(models.AttendanceSession.started_at.asc())
        .all()
    )

    target_session = None
    if session_id:
        target_session = next((s for s in sessions if s.id == session_id), None)
        if not target_session:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found for this class")
    elif sessions:
        target_session = sessions[-1]

    latest_report = _build_session_report(db, target_session, students) if target_session else None

    overall = []
    for s in students:
        attended = sum(
            1 for sess in sessions
            if db.query(models.AttendanceRecord).filter(
                models.AttendanceRecord.session_id == sess.id,
                models.AttendanceRecord.student_id == s.id,
                models.AttendanceRecord.status == "present",
            ).first()
        )
        pct = round(attended / len(sessions) * 100, 1) if sessions else 0.0
        overall.append(OverallAttendanceRow(student_id=s.id, name=s.name, reg_no=s.reg_no, sessions_attended=attended, sessions_total=len(sessions), percentage=pct))

    return ClassReport(
        class_id=school_class.id,
        class_name=school_class.name,
        class_code=school_class.code,
        sessions_run=len(sessions),
        latest_session=latest_report,
        overall=overall,
    )

@router.get("/{session_id}/export")
def export_session_excel(
    class_id: int,
    session_id: int,
    db: Session = Depends(get_db),
    teacher: models.Teacher = Depends(get_current_teacher),
):
    """Downloads one session's attendance as a real .xlsx file — class name, date,
    and every student's present/absent status with match confidence."""
    school_class = get_owned_class(class_id, db, teacher)
    session = (
        db.query(models.AttendanceSession)
        .filter(models.AttendanceSession.id == session_id, models.AttendanceSession.class_id == class_id)
        .first()
    )
    if not session:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found for this class")

    report = _build_session_report(db, session, school_class.students)

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    ws.append([f"{school_class.name} ({school_class.code}) — Section {school_class.section}"])
    ws.append([f"Session date: {session.started_at.strftime('%d %b %Y, %I:%M %p')}"])
    ws.append([f"Present: {report.present_count} / {report.total} ({report.attendance_rate}%)"])
    ws.append([])
    ws["A1"].font = Font(bold=True, size=13)

    header_row = ["Name", "Registration No.", "Status", "Match confidence (%)"]
    ws.append(header_row)
    header_idx = ws.max_row
    for cell in ws[header_idx]:
        cell.font = Font(bold=True)

    all_rows = sorted(report.present + report.absent, key=lambda r: r.name.lower())
    for row in all_rows:
        ws.append([row.name, row.reg_no, row.status.capitalize(), row.confidence if row.confidence is not None else ""])

    widths = {"A": 28, "B": 20, "C": 12, "D": 20}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"{school_class.code}_attendance_{session.started_at.strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
