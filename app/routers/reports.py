"""Attendance reporting and Excel export routes."""
import io
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import func
from sqlalchemy.orm import Session, defer
from app.database import get_db
from app.dependencies import get_current_teacher, get_owned_class
from app.schemas import ClassReport, SessionReport, RosterEntry, OverallAttendanceRow
from app import models

router = APIRouter(prefix="/classes/{class_id}/reports", tags=["reports"])

def _build_session_report(db: Session, session: models.AttendanceSession) -> SessionReport:
    records = db.query(models.AttendanceRecord).filter(models.AttendanceRecord.session_id == session.id).all()
    records_by_student = {r.student_id: r for r in records}

    all_students = (
        db.query(models.Student)
        .options(defer(models.Student.face_encoding))
        .filter(models.Student.class_id == session.class_id)
        .order_by(models.Student.name)
        .all()
    )

    present_entries, absent_entries = [], []
    for s in all_students:
        rec = records_by_student.get(s.id)
        if rec and rec.status == "present":
            present_entries.append(RosterEntry(id=s.id, name=s.name, reg_no=s.reg_no, status="present", confidence=rec.confidence))
        else:
            absent_entries.append(RosterEntry(id=s.id, name=s.name, reg_no=s.reg_no, status="absent"))

    total = len(all_students)
    p_count = len(present_entries)
    rate = round((p_count / total) * 100, 1) if total > 0 else 0.0

    return SessionReport(
        session_id=session.id,
        date=session.started_at,
        present=present_entries,
        absent=absent_entries,
        total=total,
        present_count=p_count,
        attendance_rate=rate,
    )

@router.get("", response_model=ClassReport)
def get_class_report(
    class_id: int,
    session_id: Optional[int] = None,
    db: Session = Depends(get_db),
    teacher: models.Teacher = Depends(get_current_teacher),
):
    school_class = get_owned_class(class_id, db, teacher)
    sessions = (
        db.query(models.AttendanceSession)
        .filter(models.AttendanceSession.class_id == class_id)
        .order_by(models.AttendanceSession.started_at.desc())
        .all()
    )
    sessions_run = len(sessions)

    target_session = None
    if session_id:
        target_session = next((s for s in sessions if s.id == session_id), None)
    elif sessions:
        target_session = sessions[0]

    session_report = _build_session_report(db, target_session) if target_session else None

    # Overall summary: one grouped query instead of one COUNT per student.
    students = db.query(models.Student).options(defer(models.Student.face_encoding)) \
        .filter(models.Student.class_id == class_id).order_by(models.Student.name).all()
    attended_by_student = dict(
        db.query(models.AttendanceRecord.student_id, func.count(models.AttendanceRecord.id))
        .join(models.AttendanceSession, models.AttendanceRecord.session_id == models.AttendanceSession.id)
        .filter(models.AttendanceSession.class_id == class_id, models.AttendanceRecord.status == "present")
        .group_by(models.AttendanceRecord.student_id)
        .all()
    )
    overall_rows = []
    for s in students:
        attended = attended_by_student.get(s.id, 0)
        pct = round((attended / sessions_run) * 100, 1) if sessions_run > 0 else 0.0
        overall_rows.append(
            OverallAttendanceRow(
                student_id=s.id,
                name=s.name,
                reg_no=s.reg_no,
                sessions_attended=attended,
                sessions_total=sessions_run,
                percentage=pct,
            )
        )

    return ClassReport(
        class_id=school_class.id,
        class_name=school_class.name,
        class_code=school_class.code,
        sessions_run=sessions_run,
        latest_session=session_report,
        overall=overall_rows,
    )

@router.get("/{session_id}/export")
def export_session_excel(
    class_id: int,
    session_id: int,
    db: Session = Depends(get_db),
    teacher: models.Teacher = Depends(get_current_teacher),
):
    school_class = get_owned_class(class_id, db, teacher)
    session = (
        db.query(models.AttendanceSession)
        .filter(models.AttendanceSession.id == session_id, models.AttendanceSession.class_id == class_id)
        .first()
    )
    if not session:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Session not found")

    report = _build_session_report(db, session)

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    ws.append(["FaceSync Attendance Report"])
    ws.append([f"Class: {school_class.name} ({school_class.code})"])
    ws.append([f"Session Date: {session.started_at.strftime('%Y-%m-%d %H:%M:%S')} UTC"])
    ws.append([f"Present: {report.present_count}/{report.total} ({report.attendance_rate}%)"])
    ws.append([])

    headers = ["Reg No", "Student Name", "Status", "Match Confidence"]
    ws.append(headers)

    for cell in ws[6]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1B2440", end_color="1B2440", fill_type="solid")

    for entry in report.present + report.absent:
        ws.append([
            entry.reg_no,
            entry.name,
            entry.status.upper(),
            f"{entry.confidence:.1f}%" if entry.confidence is not None else "N/A",
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"attendance_{school_class.code}_{session_id}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
